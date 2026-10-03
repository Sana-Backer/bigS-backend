"""
inventory/excel.py – stock levels out to .xlsx, stocktake counts back in.

Export columns : SKU | Name | Category | Type | On hand | Low-stock threshold | Status
Import columns : SKU | Counted quantity | Note(optional)
(Header names are matched case-insensitively; only SKU + a quantity column are required.
 You can upload an exported file after editing its "On hand" column.)
"""

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from catalog.models import Product, ProductVariant

from .exceptions import InventoryError
from .models import StockMovement
from . import services

QTY_HEADERS = ("counted quantity", "quantity", "on hand", "qty", "stock_quantity")
MAX_ROWS = 5000


def export_stock(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    headers = ["SKU", "Name", "Category", "Type", "On hand", "Low-stock threshold", "Status"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for r in rows:
        ws.append([r["sku"], r["name"], r["category"] or "", r["kind"],
                   r["stock_quantity"], r["low_stock_threshold"], r["status"]])
    for i, w in enumerate([22, 48, 22, 10, 10, 20, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_int(raw):
    if raw is None or str(raw).strip() == "":
        raise ValueError("quantity is empty")
    f = float(raw)
    if f != int(f) or f < 0:
        raise ValueError("quantity must be a whole number ≥ 0")
    return int(f)


def import_stocktake(fileobj, *, user=None, dry_run=False, note="Excel stocktake"):
    """
    Returns {"ok", "dry_run", "changes": [...], "unchanged": n, "errors": [...]}.
    Nothing is written unless every row is valid (all-or-nothing), and never on dry_run.
    """
    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception:
        raise InventoryError("Could not read the file – is it a valid .xlsx?")
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h or "").strip().lower() for h in next(rows)]
    except StopIteration:
        raise InventoryError("The file is empty.")

    if "sku" not in header:
        raise InventoryError("Missing required column: SKU.")
    qty_col = next((header.index(h) for h in QTY_HEADERS if h in header), None)
    if qty_col is None:
        raise InventoryError("Missing quantity column (use 'Counted quantity' or 'On hand').")
    sku_col = header.index("sku")
    note_col = header.index("note") if "note" in header else None

    errors, parsed, seen = [], [], set()
    for n, row in enumerate(rows, start=2):
        if n - 1 > MAX_ROWS:
            raise InventoryError(f"Too many rows (max {MAX_ROWS}).")
        if row is None or all(c in (None, "") for c in row):
            continue
        sku = str(row[sku_col] or "").strip()
        if not sku:
            errors.append({"row": n, "error": "SKU is empty"})
            continue
        if sku in seen:
            errors.append({"row": n, "sku": sku, "error": "duplicate SKU in file"})
            continue
        seen.add(sku)
        try:
            qty = _parse_int(row[qty_col])
        except (ValueError, TypeError) as e:
            errors.append({"row": n, "sku": sku, "error": str(e)})
            continue
        row_note = str(row[note_col]).strip() if note_col is not None and row[note_col] else ""
        parsed.append((n, sku, qty, row_note))

    variants = {v.sku: v for v in ProductVariant.objects.select_related("product").filter(sku__in=seen)}
    products = {p.sku: p for p in Product.objects.filter(sku__in=seen)}
    plan = []
    for n, sku, qty, row_note in parsed:
        target = variants.get(sku)
        if target is None:
            target = products.get(sku)
            if target is not None and target.variants.exists():
                errors.append({"row": n, "sku": sku, "error": "product has variants – use a variant SKU"})
                continue
        if target is None:
            errors.append({"row": n, "sku": sku, "error": "SKU not found"})
            continue
        plan.append((n, sku, target, qty, row_note))

    changes, unchanged = [], 0
    for n, sku, target, qty, row_note in plan:
        if target.stock_quantity == qty:
            unchanged += 1
        else:
            changes.append({"row": n, "sku": sku, "from": target.stock_quantity, "to": qty})

    result = {"ok": not errors, "dry_run": dry_run, "changes": changes,
              "unchanged": unchanged, "errors": errors}
    if errors or dry_run:
        return result

    from django.db import transaction
    with transaction.atomic():
        for n, sku, target, qty, row_note in plan:
            services.set_stock(target, qty, user=user, note=row_note or note,
                               movement_type=StockMovement.Type.STOCKTAKE)
    return result
