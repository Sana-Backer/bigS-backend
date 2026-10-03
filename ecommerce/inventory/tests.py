"""
inventory/tests.py  –  python manage.py test inventory
"""

import uuid
from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from catalog.models import Category, Product, ProductVariant
from users.models import User

from . import services
from .exceptions import InsufficientStockError, InventoryError
from .models import PurchaseOrder, StockMovement, StockPolicy, Supplier

M = StockMovement.Type
BASE = "/api/admin/inventory/"


def sku():
    return f"SKU-{uuid.uuid4().hex[:8].upper()}"


class Base(TestCase):
    def setUp(self):
        self.cat = Category.objects.create(name="Skin Care")
        self.plain = Product.objects.create(category=self.cat, name="Soap", sku=sku(),
                                            base_price=Decimal("100"), stock_quantity=50)
        self.parent = Product.objects.create(category=self.cat, name="Serum", sku=sku(),
                                             base_price=Decimal("500"))
        self.v30 = ProductVariant.objects.create(product=self.parent, name="30ml", sku=sku(),
                                                 attributes={"size": "30ml"}, price=Decimal("500"),
                                                 stock_quantity=20)
        self.v50 = ProductVariant.objects.create(product=self.parent, name="50ml", sku=sku(),
                                                 attributes={"size": "50ml"}, price=Decimal("800"),
                                                 stock_quantity=0)
        StockMovement.objects.all().delete()  # start each test with an empty ledger

        self.customer = User.objects.create_user(email="c@example.com", password="pass12345")
        self.staff = User.objects.create_staff(email="s@example.com", password="pass12345")
        self.manager = User.objects.create_user(email="m@example.com", password="pass12345",
                                                role=User.Role.MANAGER)
        self.client = APIClient()

    def as_user(self, user):
        self.client.force_authenticate(user)
        return self.client


class ServiceTests(Base):
    def test_movement_updates_stock_and_ledger(self):
        mv = services.record_movement(self.v30, -5, M.SALE, reference="ORD-1", user=self.staff)
        self.v30.refresh_from_db()
        self.assertEqual(self.v30.stock_quantity, 15)
        self.assertEqual((mv.quantity_before, mv.quantity_after, mv.quantity_change), (20, 15, -5))
        self.assertEqual(StockMovement.objects.count(), 1)  # signal must not double-log

    def test_cannot_go_negative(self):
        with self.assertRaises(InsufficientStockError):
            services.record_movement(self.v30, -21, M.ADJUSTMENT)
        self.v30.refresh_from_db()
        self.assertEqual(self.v30.stock_quantity, 20)
        self.assertEqual(StockMovement.objects.count(), 0)

    def test_variant_parent_product_is_rejected(self):
        with self.assertRaises(InventoryError):
            services.resolve_target(product_id=self.parent.id)

    def test_set_stock_noop_and_change(self):
        self.assertIsNone(services.set_stock(self.plain, 50))
        mv = services.set_stock(self.plain, 42, note="count")
        self.assertEqual(mv.quantity_change, -8)

    def test_external_save_is_logged(self):
        self.v30.stock_quantity = 25
        self.v30.save()
        mv = StockMovement.objects.get()
        self.assertEqual((mv.movement_type, mv.quantity_change), (M.EXTERNAL, 5))

    def test_new_product_with_opening_stock_logs_initial(self):
        p = Product.objects.create(category=self.cat, name="New", sku=sku(),
                                   base_price=Decimal("10"), stock_quantity=7)
        mv = StockMovement.objects.get(product=p)
        self.assertEqual((mv.movement_type, mv.quantity_after), (M.INITIAL, 7))


class StockApiTests(Base):
    def test_permissions(self):
        self.assertEqual(self.client.get(BASE + "stock/").status_code, 401)
        self.assertEqual(self.as_user(self.customer).get(BASE + "stock/").status_code, 403)
        self.assertEqual(self.as_user(self.staff).get(BASE + "stock/").status_code, 200)

    def test_list_filters_and_policy_threshold(self):
        c = self.as_user(self.staff)
        data = c.get(BASE + "stock/").json()
        self.assertEqual(data["count"], 3)  # soap + 2 variants (parent product not listed)
        out = c.get(BASE + "stock/?status=out").json()["data"]
        self.assertEqual([r["sku"] for r in out], [self.v50.sku])
        self.assertEqual(c.get(BASE + "stock/?status=low").json()["count"], 0)
        StockPolicy.objects.create(variant=self.v30, low_stock_threshold=25)
        low = c.get(BASE + "stock/?status=low").json()["data"]
        self.assertEqual([r["sku"] for r in low], [self.v30.sku])
        self.assertEqual(c.get(BASE + "stock/?status=attention").json()["count"], 2)
        self.assertEqual(c.get(BASE + "stock/?q=soap").json()["count"], 1)

    def test_summary(self):
        d = self.as_user(self.staff).get(BASE + "stock/summary/").json()["data"]
        self.assertEqual(d["units_on_hand"], 70)
        self.assertEqual(d["out_of_stock"], 1)

    def test_adjust(self):
        c = self.as_user(self.staff)
        r = c.post(BASE + "stock/adjust/", {"variant_id": str(self.v30.id), "change": 10,
                                            "movement_type": "return", "note": "RMA"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.v30.refresh_from_db()
        self.assertEqual(self.v30.stock_quantity, 30)
        self.assertEqual(r.json()["data"]["created_by_email"], "s@example.com")

    def test_adjust_validation(self):
        c = self.as_user(self.staff)
        r = c.post(BASE + "stock/adjust/", {"variant_id": str(self.v30.id), "change": -99}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["status"], "error")
        r = c.post(BASE + "stock/adjust/", {"product_id": str(self.parent.id), "change": 1}, format="json")
        self.assertEqual(r.status_code, 400)  # parent has variants
        r = c.post(BASE + "stock/adjust/", {"change": 1}, format="json")
        self.assertEqual(r.status_code, 400)
        r = c.post(BASE + "stock/adjust/", {"variant_id": str(self.v30.id), "change": 1,
                                            "movement_type": "sale"}, format="json")
        self.assertEqual(r.status_code, 400)  # system-only type

    def test_bulk_adjust_is_atomic(self):
        c = self.as_user(self.staff)
        payload = {"items": [
            {"variant_id": str(self.v30.id), "change": 5},
            {"product_id": str(self.plain.id), "change": -500},   # fails
        ]}
        r = c.post(BASE + "stock/bulk-adjust/", payload, format="json")
        self.assertEqual(r.status_code, 400)
        self.v30.refresh_from_db()
        self.assertEqual(self.v30.stock_quantity, 20)
        self.assertEqual(StockMovement.objects.count(), 0)

    def test_set_requires_manager(self):
        body = {"variant_id": str(self.v30.id), "quantity": 12}
        self.assertEqual(self.as_user(self.staff).post(BASE + "stock/set/", body, format="json").status_code, 403)
        r = self.as_user(self.manager).post(BASE + "stock/set/", body, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.v30.refresh_from_db()
        self.assertEqual(self.v30.stock_quantity, 12)

    def test_movements_filter(self):
        services.record_movement(self.v30, -1, M.SALE, reference="ORD-9")
        services.record_movement(self.plain, 3, M.RETURN)
        c = self.as_user(self.staff)
        self.assertEqual(c.get(BASE + "movements/").json()["count"], 2)
        self.assertEqual(c.get(BASE + "movements/?reference=ord-9").json()["count"], 1)
        self.assertEqual(c.get(BASE + f"movements/?sku={self.plain.sku}").json()["count"], 1)
        self.assertEqual(c.get(BASE + "movements/?type=sale").json()["count"], 1)
        self.assertEqual(c.get(BASE + "movements/?date_from=bad").status_code, 400)


def xlsx(rows, header=("SKU", "Counted quantity", "Note")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    buf = BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("count.xlsx", buf.getvalue(),
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class ExcelTests(Base):
    def test_export(self):
        r = self.as_user(self.staff).get(BASE + "stock/export/")
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[1]][:2], ["SKU", "Name"])
        self.assertEqual(ws.max_row, 4)

    def test_import_apply_and_dry_run(self):
        c = self.as_user(self.manager)
        rows = [(self.v30.sku, 12, "shelf A"), (self.plain.sku, 50, ""), (self.v50.sku, 4, "")]
        r = c.post(BASE + "stock/import/?dry_run=true", {"file": xlsx(rows)}, format="multipart")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(len(r.json()["data"]["changes"]), 2)
        self.assertEqual(r.json()["data"]["unchanged"], 1)
        self.v30.refresh_from_db()
        self.assertEqual(self.v30.stock_quantity, 20)  # dry run changed nothing

        r = c.post(BASE + "stock/import/", {"file": xlsx(rows)}, format="multipart")
        self.assertEqual(r.status_code, 200, r.content)
        self.v30.refresh_from_db(); self.v50.refresh_from_db()
        self.assertEqual((self.v30.stock_quantity, self.v50.stock_quantity), (12, 4))
        mv = StockMovement.objects.get(variant=self.v30)
        self.assertEqual((mv.movement_type, mv.note), (M.STOCKTAKE, "shelf A"))
        self.assertEqual(StockMovement.objects.count(), 2)

    def test_import_is_all_or_nothing(self):
        rows = [(self.v30.sku, 1, ""), ("NOPE", 3, ""), (self.plain.sku, -2, ""), (self.parent.sku, 5, "")]
        r = self.as_user(self.manager).post(BASE + "stock/import/", {"file": xlsx(rows)}, format="multipart")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(len(r.json()["errors"]), 3)
        self.v30.refresh_from_db()
        self.assertEqual(self.v30.stock_quantity, 20)

    def test_import_needs_manager_and_xlsx(self):
        self.assertEqual(self.as_user(self.staff).post(BASE + "stock/import/", {"file": xlsx([])},
                                                       format="multipart").status_code, 403)
        bad = SimpleUploadedFile("x.csv", b"a,b")
        self.assertEqual(self.as_user(self.manager).post(BASE + "stock/import/", {"file": bad},
                                                         format="multipart").status_code, 400)


class PurchaseOrderTests(Base):
    def setUp(self):
        super().setUp()
        self.supplier = Supplier.objects.create(name="Acme Beauty Wholesale")

    def create_po(self):
        body = {"supplier": str(self.supplier.id), "items": [
            {"variant_id": str(self.v50.id), "quantity_ordered": 30, "unit_cost": "250.00"},
            {"product_id": str(self.plain.id), "quantity_ordered": 10, "unit_cost": "40.00"},
        ]}
        r = self.as_user(self.manager).post(BASE + "purchase-orders/", body, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        return r.json()

    def test_lifecycle(self):
        po = self.create_po()
        pid = po["id"]
        self.assertEqual(po["status"], "draft")
        self.assertEqual(Decimal(po["total_cost"]), Decimal("7900.00"))

        c = self.as_user(self.staff)
        # cannot receive a draft
        self.assertEqual(c.post(f"{BASE}purchase-orders/{pid}/receive/", {}, format="json").status_code, 400)
        # staff cannot place the order
        self.assertEqual(c.post(f"{BASE}purchase-orders/{pid}/mark-ordered/").status_code, 403)

        self.assertEqual(self.as_user(self.manager).post(f"{BASE}purchase-orders/{pid}/mark-ordered/").status_code, 200)

        line = next(i for i in po["items"] if i["sku"] == self.v50.sku)
        c = self.as_user(self.staff)
        r = c.post(f"{BASE}purchase-orders/{pid}/receive/",
                   {"lines": [{"item_id": line["id"], "quantity": 20}]}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["data"]["status"], "partially_received")
        self.v50.refresh_from_db()
        self.assertEqual(self.v50.stock_quantity, 20)

        # over-receiving is rejected and changes nothing
        r = c.post(f"{BASE}purchase-orders/{pid}/receive/",
                   {"lines": [{"item_id": line["id"], "quantity": 11}]}, format="json")
        self.assertEqual(r.status_code, 400)

        # receive the rest
        r = c.post(f"{BASE}purchase-orders/{pid}/receive/", {}, format="json")
        self.assertEqual(r.json()["data"]["status"], "received")
        self.v50.refresh_from_db(); self.plain.refresh_from_db()
        self.assertEqual((self.v50.stock_quantity, self.plain.stock_quantity), (30, 60))
        refs = set(StockMovement.objects.filter(movement_type=M.PURCHASE_RECEIPT)
                   .values_list("reference", flat=True))
        self.assertEqual(refs, {po["po_number"]})
        self.assertEqual(StockMovement.objects.count(), 3)

    def test_cancel_and_edit_rules(self):
        po = self.create_po()
        m = self.as_user(self.manager)
        self.assertEqual(m.post(f"{BASE}purchase-orders/{po['id']}/cancel/").status_code, 200)
        self.assertEqual(PurchaseOrder.objects.get(pk=po["id"]).status, "cancelled")
        self.assertEqual(m.post(f"{BASE}purchase-orders/{po['id']}/cancel/").status_code, 400)

    def test_po_validation(self):
        m = self.as_user(self.manager)
        r = m.post(BASE + "purchase-orders/", {"supplier": str(self.supplier.id), "items": []}, format="json")
        self.assertEqual(r.status_code, 400)
        r = m.post(BASE + "purchase-orders/", {"supplier": str(self.supplier.id), "items": [
            {"product_id": str(self.parent.id), "quantity_ordered": 5}]}, format="json")
        self.assertEqual(r.status_code, 400)  # parent has variants

    def test_supplier_with_history_is_deactivated_not_deleted(self):
        self.create_po()
        r = self.as_user(self.manager).delete(f"{BASE}suppliers/{self.supplier.id}/")
        self.assertEqual(r.status_code, 204)
        self.supplier.refresh_from_db()
        self.assertFalse(self.supplier.is_active)


class OrderIntegrationTests(Base):
    """Checkout writes a 'sale' ledger row; cancelling writes 'order_cancel'."""

    def test_checkout_and_cancel_hit_the_ledger(self):
        from cart.models import Cart, CartItem
        from orders import services as order_services

        cart = Cart.objects.create(user=self.customer)
        CartItem.objects.create(cart=cart, product=self.parent, variant=self.v30, quantity=3,
                                unit_price=self.v30.price)
        CartItem.objects.create(cart=cart, product=self.plain, quantity=2, unit_price=self.plain.base_price)
        addr = {"full_name": "A B", "phone": "9999999999", "line1": "1 Street", "city": "Kochi",
                "state": "KL", "postal_code": "682001", "country": "IN"}
        order = order_services.create_order_from_cart(
            user=self.customer, cart=cart, billing_address=addr, shipping_address=addr,
            customer_email="c@example.com")

        self.v30.refresh_from_db(); self.plain.refresh_from_db()
        self.assertEqual((self.v30.stock_quantity, self.plain.stock_quantity), (17, 48))
        sales = StockMovement.objects.filter(movement_type=M.SALE, reference=order.order_number)
        self.assertEqual(sales.count(), 2)
        self.assertEqual(sum(s.quantity_change for s in sales), -5)

        order_services.cancel_order(order, changed_by=self.customer)
        self.v30.refresh_from_db(); self.plain.refresh_from_db()
        self.assertEqual((self.v30.stock_quantity, self.plain.stock_quantity), (20, 50))
        self.assertEqual(StockMovement.objects.filter(movement_type=M.ORDER_CANCEL,
                                                      reference=order.order_number).count(), 2)
        self.assertFalse(StockMovement.objects.filter(movement_type=M.EXTERNAL).exists())
